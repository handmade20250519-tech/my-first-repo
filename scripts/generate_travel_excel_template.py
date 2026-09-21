#!/usr/bin/env python3
"""旅行計画用Excelテンプレート生成スクリプト。

8シート(旅行概要/日別スケジュール/予約管理/旅行予算/持ち物チェックリスト/
行きたい場所・食べたいもの/お土産リスト/緊急連絡先)から成る .xlsx を出力する。
"""

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import CellIsRule, Rule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

OUTPUT_PATH = "旅行計画テンプレート.xlsx"

# ---------------------------------------------------------------------------
# 配色・フォント定義(派手すぎない旅行らしい配色)
# ---------------------------------------------------------------------------
COLOR_TITLE_BG = "1F4E4A"      # 濃いティール(タイトルバナー)
COLOR_HEADER_BG = "3D7B75"     # ティール(表ヘッダー)
COLOR_HEADER_FG = "FFFFFF"
COLOR_SUMMARY_BG = "9FC6BE"    # 淡いティール(サマリー見出し)
COLOR_INPUT_BG = "DCEEF3"      # 淡い水色 = 入力欄
COLOR_CALC_BG = "FCEBC8"       # 淡い砂色 = 自動計算欄
COLOR_EXAMPLE_FG = "8C8C8C"    # 入力例の文字色(グレー)
COLOR_BAND_BG = "F6F4EE"       # 薄いクリーム(縞模様)
COLOR_BORDER = "CFCABB"
COLOR_ACCENT = "C97B4A"        # テラコッタ(アクセント)
COLOR_GOOD = "2E7D46"
COLOR_WARN = "B4530A"
COLOR_BAD = "B03A2E"

FONT_NAME = "游ゴシック"

F_TITLE = Font(name=FONT_NAME, size=16, bold=True, color=COLOR_HEADER_FG)
F_LEGEND = Font(name=FONT_NAME, size=9, italic=True, color="5B5B5B")
F_HEADER = Font(name=FONT_NAME, size=10, bold=True, color=COLOR_HEADER_FG)
F_LABEL = Font(name=FONT_NAME, size=10, bold=True, color="1F4E4A")
F_BODY = Font(name=FONT_NAME, size=10, color="333333")
F_EXAMPLE = Font(name=FONT_NAME, size=10, italic=True, color=COLOR_EXAMPLE_FG)
F_CALC = Font(name=FONT_NAME, size=10, color="7A4A16")
F_TOTAL_LABEL = Font(name=FONT_NAME, size=10, bold=True, color="1F4E4A")
F_SUMMARY_HEAD = Font(name=FONT_NAME, size=10, bold=True, color=COLOR_HEADER_FG)
F_SUMMARY_BIG = Font(name=FONT_NAME, size=13, bold=True, color="1F4E4A")
F_NOTE = Font(name=FONT_NAME, size=9, italic=True, color="8C8C8C")

FILL_TITLE = PatternFill("solid", fgColor=COLOR_TITLE_BG)
FILL_HEADER = PatternFill("solid", fgColor=COLOR_HEADER_BG)
FILL_SUMMARY = PatternFill("solid", fgColor=COLOR_SUMMARY_BG)
FILL_INPUT = PatternFill("solid", fgColor=COLOR_INPUT_BG)
FILL_CALC = PatternFill("solid", fgColor=COLOR_CALC_BG)
FILL_BAND = PatternFill("solid", fgColor=COLOR_BAND_BG)
FILL_WHITE = PatternFill("solid", fgColor="FFFFFF")

thin = Side(style="thin", color=COLOR_BORDER)
BORDER_ALL = Border(left=thin, right=thin, top=thin, bottom=thin)
thick_top = Side(style="thin", color="1F4E4A")
BORDER_TOTAL = Border(left=thin, right=thin, top=thick_top, bottom=thin)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
ALIGN_LEFT_NOWRAP = Alignment(horizontal="left", vertical="center", indent=1)
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center", indent=1)

CUR_FMT = '"¥"#,##0;[Red]-"¥"#,##0'
CUR_FMT_SIMPLE = '"¥"#,##0'
DATE_FMT = 'yyyy"年"m"月"d"日"'
DATE_W_FMT = 'm"月"d"日"'
TIME_FMT = "hh:mm"
PCT_FMT = "0.0%"


def set_cell(ws, coord, value=None, font=F_BODY, fill=None, align=ALIGN_LEFT,
             border=BORDER_ALL, number_format=None):
    cell = ws[coord]
    if value is not None:
        cell.value = value
    cell.font = font
    if fill is not None:
        cell.fill = fill
    if align is not None:
        cell.alignment = align
    if border is not None:
        cell.border = border
    if number_format is not None:
        cell.number_format = number_format
    return cell


def set_col_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def title_banner(ws, last_col_letter, text, title_row_h=30, legend_row_h=16):
    ws.merge_cells(f"A1:{last_col_letter}1")
    set_cell(ws, "A1", text, font=F_TITLE, fill=FILL_TITLE,
             align=Alignment(horizontal="left", vertical="center", indent=1), border=None)
    ws.row_dimensions[1].height = title_row_h

    ws.merge_cells(f"A2:{last_col_letter}2")
    set_cell(ws, "A2", "🔵 水色のセル = 入力欄です。自由に書き換えてください　　🟠 オレンジのセル = 自動計算欄です(編集不要)",
              font=F_LEGEND, fill=FILL_WHITE, align=Alignment(horizontal="left", vertical="center", indent=1),
              border=None)
    ws.row_dimensions[2].height = legend_row_h


def table_header(ws, row, headers_by_col):
    for col, text in headers_by_col.items():
        set_cell(ws, f"{col}{row}", text, font=F_HEADER, fill=FILL_HEADER,
                  align=ALIGN_CENTER, border=BORDER_ALL)
    ws.row_dimensions[row].height = 20


def band_row(ws, row, cols, fill_even):
    if fill_even:
        for col in cols:
            ws[f"{col}{row}"].fill = FILL_BAND


def page_setup(ws, orientation="portrait", fit_width=1, repeat_rows="1:3"):
    ws.page_setup.orientation = orientation
    ws.page_setup.fitToWidth = fit_width
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.print_title_rows = repeat_rows
    ws.sheet_view.showGridLines = False


def add_list_validation(ws, cell_range, options, allow_blank=True):
    dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=allow_blank,
                         showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(cell_range)
    return dv


def equal_rule(color_fg, bold=False):
    return Font(name=FONT_NAME, size=10, color=color_fg, bold=bold)


wb = Workbook()
wb.remove(wb.active)

# ===========================================================================
# 1. 旅行概要
# ===========================================================================
ws = wb.create_sheet("旅行概要")
set_col_widths(ws, {"A": 16, "B": 12, "C": 12, "D": 12, "E": 12, "F": 12, "G": 12})
title_banner(ws, "G", "✈ 旅行概要")

fields = [
    (4, "旅行名", "沖縄家族旅行2026", False),
    (5, "目的地", "沖縄県 那覇市・恩納村", False),
    (6, "出発日", None, False),
    (7, "帰宅日", None, False),
    (8, "旅行日数(自動計算)", None, True),
    (9, "参加者", "父・母・長女・長男", False),
    (10, "交通手段", "飛行機(ANA123便) + レンタカー", False),
    (11, "宿泊先", "〇〇リゾート&スパ 那覇", False),
]
import datetime
for row, label, example, is_calc in fields:
    set_cell(ws, f"A{row}", label, font=F_LABEL, fill=FILL_SUMMARY, align=ALIGN_LEFT_NOWRAP)
    ws.merge_cells(f"B{row}:G{row}")
    if is_calc:
        set_cell(ws, f"B{row}", '=IF(AND(B7<>"",B6<>""),(B7-B6)&"泊"&(B7-B6+1)&"日","")',
                  font=F_CALC, fill=FILL_CALC, align=ALIGN_LEFT_NOWRAP)
    elif label == "出発日":
        set_cell(ws, f"B{row}", datetime.date(2026, 5, 3), font=F_BODY, fill=FILL_INPUT,
                  align=ALIGN_LEFT_NOWRAP, number_format=DATE_FMT)
    elif label == "帰宅日":
        set_cell(ws, f"B{row}", datetime.date(2026, 5, 5), font=F_BODY, fill=FILL_INPUT,
                  align=ALIGN_LEFT_NOWRAP, number_format=DATE_FMT)
    else:
        set_cell(ws, f"B{row}", example, font=F_BODY, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP)
    ws.row_dimensions[row].height = 22

set_cell(ws, "A13", "【凡例】", font=F_LABEL, fill=None, align=ALIGN_LEFT_NOWRAP, border=None)
set_cell(ws, "A14", "", fill=FILL_INPUT, border=BORDER_ALL)
ws.merge_cells("B14:G14")
set_cell(ws, "B14", "＝入力欄です。旅行の情報を自由に書き換えてください", font=F_BODY, fill=None,
          align=ALIGN_LEFT_NOWRAP, border=None)
set_cell(ws, "A15", "", fill=FILL_CALC, border=BORDER_ALL)
ws.merge_cells("B15:G15")
set_cell(ws, "B15", "＝自動計算欄です。数式が入っているため編集不要です", font=F_BODY, fill=None,
          align=ALIGN_LEFT_NOWRAP, border=None)

page_setup(ws, orientation="portrait", repeat_rows="1:2")

# ===========================================================================
# 2. 日別スケジュール
# ===========================================================================
ws = wb.create_sheet("日別スケジュール")
set_col_widths(ws, {"A": 13, "B": 8, "C": 22, "D": 18, "E": 12, "F": 10, "G": 10, "H": 24})
title_banner(ws, "H", "📅 日別スケジュール")
table_header(ws, 3, {"A": "日付", "B": "時間", "C": "予定", "D": "場所", "E": "移動手段",
                       "F": "予約有無", "G": "費用", "H": "メモ"})
ws.freeze_panes = "A4"

examples = [
    (datetime.date(2026, 5, 3), datetime.time(9, 0), "空港集合・チェックイン", "羽田空港 第2ターミナル",
     "電車", "済", 0, "1時間前集合"),
    (datetime.date(2026, 5, 3), datetime.time(11, 0), "那覇空港到着・レンタカー受取", "那覇空港", "飛行機",
     "済", 15000, "ANA123便"),
    (datetime.date(2026, 5, 3), datetime.time(15, 0), "ホテルチェックイン", "〇〇リゾート&スパ", "レンタカー",
     "済", 0, ""),
]
DATA_START = 4
DATA_END = 28
for i, (d, t, plan, place, move, res, cost, memo) in enumerate(examples):
    r = DATA_START + i
    set_cell(ws, f"A{r}", d, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP, number_format=DATE_W_FMT)
    set_cell(ws, f"B{r}", t, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_CENTER, number_format=TIME_FMT)
    set_cell(ws, f"C{r}", plan, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP)
    set_cell(ws, f"D{r}", place, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP)
    set_cell(ws, f"E{r}", move, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_CENTER)
    set_cell(ws, f"F{r}", res, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_CENTER)
    set_cell(ws, f"G{r}", cost, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_RIGHT, number_format=CUR_FMT_SIMPLE)
    set_cell(ws, f"H{r}", memo, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP)
    ws.row_dimensions[r].height = 20

for r in range(DATA_START + len(examples), DATA_END + 1):
    for col, fmt in (("A", DATE_W_FMT), ("B", TIME_FMT), ("C", None), ("D", None), ("E", None),
                      ("F", None), ("G", CUR_FMT_SIMPLE), ("H", None)):
        align = ALIGN_CENTER if col in ("A", "B", "E", "F") else (ALIGN_RIGHT if col == "G" else ALIGN_LEFT_NOWRAP)
        set_cell(ws, f"{col}{r}", None, font=F_BODY, fill=FILL_INPUT, align=align, number_format=fmt)
    ws.row_dimensions[r].height = 20
    band_row(ws, r, ["A", "B", "C", "D", "E", "F", "G", "H"], (r - DATA_START) % 2 == 1)

TOTAL_ROW = DATA_END + 1
ws.merge_cells(f"E{TOTAL_ROW}:F{TOTAL_ROW}")
set_cell(ws, f"E{TOTAL_ROW}", "合計費用", font=F_TOTAL_LABEL, fill=FILL_SUMMARY, align=ALIGN_RIGHT, border=BORDER_TOTAL)
set_cell(ws, f"G{TOTAL_ROW}", f"=SUM(G{DATA_START}:G{DATA_END})", font=F_CALC, fill=FILL_CALC,
          align=ALIGN_RIGHT, border=BORDER_TOTAL, number_format=CUR_FMT)
set_cell(ws, f"H{TOTAL_ROW}", "", fill=FILL_SUMMARY, border=BORDER_TOTAL)
set_cell(ws, f"A{TOTAL_ROW}", "", fill=FILL_SUMMARY, border=BORDER_TOTAL)
set_cell(ws, f"B{TOTAL_ROW}", "", fill=FILL_SUMMARY, border=BORDER_TOTAL)
set_cell(ws, f"C{TOTAL_ROW}", "", fill=FILL_SUMMARY, border=BORDER_TOTAL)
set_cell(ws, f"D{TOTAL_ROW}", "", fill=FILL_SUMMARY, border=BORDER_TOTAL)

add_list_validation(ws, f"F{DATA_START}:F{DATA_END}", "済,未定,不要")
ws.conditional_formatting.add(f"F{DATA_START}:F{DATA_END}",
    CellIsRule(operator="equal", formula=['"済"'], font=equal_rule(COLOR_GOOD, True)))
ws.conditional_formatting.add(f"F{DATA_START}:F{DATA_END}",
    CellIsRule(operator="equal", formula=['"未定"'], font=equal_rule(COLOR_WARN, True)))

page_setup(ws, orientation="landscape")

# ===========================================================================
# 3. 予約管理
# ===========================================================================
ws = wb.create_sheet("予約管理")
set_col_widths(ws, {"A": 12, "B": 20, "C": 16, "D": 14, "E": 11, "F": 10, "G": 16, "H": 24, "I": 20})
title_banner(ws, "I", "📝 予約管理")
table_header(ws, 3, {"A": "カテゴリ", "B": "予約先", "C": "予約日時", "D": "予約番号", "E": "料金",
                       "F": "支払済み", "G": "連絡先", "H": "URL", "I": "メモ"})
ws.freeze_panes = "A4"

res_examples = [
    ("交通", "ANA(全日空)", datetime.datetime(2026, 3, 1, 10, 0), "ABC1234", 45000, "済",
     "0570-029-333", "https://www.ana.co.jp", "往復航空券"),
    ("宿泊", "〇〇リゾート&スパ 那覇", datetime.datetime(2026, 3, 1, 12, 0), "HTL-9821", 68000, "未払い",
     "098-000-0000", "https://example-resort.jp", "2泊 朝食付き"),
    ("レンタカー", "沖縄レンタカー", datetime.datetime(2026, 3, 5, 9, 0), "RC-5567", 12000, "未払い",
     "098-111-2222", "https://example-rentacar.jp", "コンパクトカー 3日間"),
]
DATA_START = 4
DATA_END = 23
for i, (cat, dest, dt, num, fee, paid, contact, url, memo) in enumerate(res_examples):
    r = DATA_START + i
    set_cell(ws, f"A{r}", cat, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_CENTER)
    set_cell(ws, f"B{r}", dest, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP)
    set_cell(ws, f"C{r}", dt, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_CENTER, number_format="yyyy/m/d hh:mm")
    set_cell(ws, f"D{r}", num, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_CENTER)
    set_cell(ws, f"E{r}", fee, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_RIGHT, number_format=CUR_FMT_SIMPLE)
    set_cell(ws, f"F{r}", paid, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_CENTER)
    set_cell(ws, f"G{r}", contact, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP)
    set_cell(ws, f"H{r}", url, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP)
    set_cell(ws, f"I{r}", memo, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP)
    ws.row_dimensions[r].height = 20

for r in range(DATA_START + len(res_examples), DATA_END + 1):
    for col, fmt, align in (("A", None, ALIGN_CENTER), ("B", None, ALIGN_LEFT_NOWRAP),
                              ("C", "yyyy/m/d hh:mm", ALIGN_CENTER), ("D", None, ALIGN_CENTER),
                              ("E", CUR_FMT_SIMPLE, ALIGN_RIGHT), ("F", None, ALIGN_CENTER),
                              ("G", None, ALIGN_LEFT_NOWRAP), ("H", None, ALIGN_LEFT_NOWRAP),
                              ("I", None, ALIGN_LEFT_NOWRAP)):
        set_cell(ws, f"{col}{r}", None, font=F_BODY, fill=FILL_INPUT, align=align, number_format=fmt)
    ws.row_dimensions[r].height = 20
    band_row(ws, r, ["A", "B", "C", "D", "E", "F", "G", "H", "I"], (r - DATA_START) % 2 == 1)

TOTAL_ROW = DATA_END + 1
ws.merge_cells(f"C{TOTAL_ROW}:D{TOTAL_ROW}")
set_cell(ws, f"C{TOTAL_ROW}", "合計料金", font=F_TOTAL_LABEL, fill=FILL_SUMMARY, align=ALIGN_RIGHT, border=BORDER_TOTAL)
set_cell(ws, f"E{TOTAL_ROW}", f"=SUM(E{DATA_START}:E{DATA_END})", font=F_CALC, fill=FILL_CALC,
          align=ALIGN_RIGHT, border=BORDER_TOTAL, number_format=CUR_FMT)
for col in ("A", "B", "F", "G", "H", "I"):
    set_cell(ws, f"{col}{TOTAL_ROW}", "", fill=FILL_SUMMARY, border=BORDER_TOTAL)

add_list_validation(ws, f"A{DATA_START}:A{DATA_END}", "交通,宿泊,レンタカー,アクティビティ,レストラン,保険,その他")
add_list_validation(ws, f"F{DATA_START}:F{DATA_END}", "済,未払い")
ws.conditional_formatting.add(f"F{DATA_START}:F{DATA_END}",
    CellIsRule(operator="equal", formula=['"済"'], font=equal_rule(COLOR_GOOD, True)))
ws.conditional_formatting.add(f"F{DATA_START}:F{DATA_END}",
    CellIsRule(operator="equal", formula=['"未払い"'], font=equal_rule(COLOR_BAD, True)))

page_setup(ws, orientation="landscape")

# ===========================================================================
# 4. 旅行予算
# ===========================================================================
ws = wb.create_sheet("旅行予算")
set_col_widths(ws, {"A": 14, "B": 13, "C": 13, "D": 13, "E": 3, "F": 14, "G": 13, "H": 13})
title_banner(ws, "H", "💴 旅行予算")

table_header(ws, 3, {"A": "費目", "B": "予算", "C": "実績", "D": "差額"})
categories = ["交通費", "宿泊費", "食費", "観光費", "お土産", "その他"]
sample_budget = [40000, 68000, 20000, 15000, 10000, 5000]
sample_actual = [45000, 68000, 0, 0, 0, 0]
CAT_START = 4
CAT_END = CAT_START + len(categories) - 1
for i, cat in enumerate(categories):
    r = CAT_START + i
    set_cell(ws, f"A{r}", cat, font=F_BODY, fill=FILL_SUMMARY, align=ALIGN_LEFT_NOWRAP)
    b_font = F_EXAMPLE if i < 2 else F_BODY
    set_cell(ws, f"B{r}", sample_budget[i] if i < 2 else None, font=b_font, fill=FILL_INPUT,
              align=ALIGN_RIGHT, number_format=CUR_FMT_SIMPLE)
    set_cell(ws, f"C{r}", sample_actual[i] if i < 2 else None, font=b_font, fill=FILL_INPUT,
              align=ALIGN_RIGHT, number_format=CUR_FMT_SIMPLE)
    set_cell(ws, f"D{r}", f"=B{r}-C{r}", font=F_CALC, fill=FILL_CALC, align=ALIGN_RIGHT, number_format=CUR_FMT)
    ws.row_dimensions[r].height = 20

TOTAL_ROW = CAT_END + 1
set_cell(ws, f"A{TOTAL_ROW}", "合計", font=F_TOTAL_LABEL, fill=FILL_SUMMARY, align=ALIGN_LEFT_NOWRAP, border=BORDER_TOTAL)
set_cell(ws, f"B{TOTAL_ROW}", f"=SUM(B{CAT_START}:B{CAT_END})", font=F_CALC, fill=FILL_CALC,
          align=ALIGN_RIGHT, border=BORDER_TOTAL, number_format=CUR_FMT)
set_cell(ws, f"C{TOTAL_ROW}", f"=SUM(C{CAT_START}:C{CAT_END})", font=F_CALC, fill=FILL_CALC,
          align=ALIGN_RIGHT, border=BORDER_TOTAL, number_format=CUR_FMT)
set_cell(ws, f"D{TOTAL_ROW}", f"=B{TOTAL_ROW}-C{TOTAL_ROW}", font=F_CALC, fill=FILL_CALC,
          align=ALIGN_RIGHT, border=BORDER_TOTAL, number_format=CUR_FMT)

NOTE_ROW = TOTAL_ROW + 2
ws.merge_cells(f"A{NOTE_ROW}:D{NOTE_ROW}")
set_cell(ws, f"A{NOTE_ROW}", "※「予算」と「実績」の水色セルに入力してください。「差額」は自動計算されます。", font=F_NOTE,
          fill=None, align=ALIGN_LEFT_NOWRAP, border=None)

# サマリーカード(右側)
ws.merge_cells("F3:H3")
set_cell(ws, "F3", "サマリー", font=F_SUMMARY_HEAD, fill=FILL_HEADER, align=ALIGN_CENTER, border=BORDER_ALL)
ws.row_dimensions[3].height = 20

summary_rows = [
    (4, "合計予算", f"=B{TOTAL_ROW}", CUR_FMT),
    (5, "合計実績", f"=C{TOTAL_ROW}", CUR_FMT),
    (6, "残り予算", f"=B{TOTAL_ROW}-C{TOTAL_ROW}", CUR_FMT),
    (7, "予算消化率", f'=IFERROR(C{TOTAL_ROW}/B{TOTAL_ROW},0)', PCT_FMT),
]
for r, label, formula, fmt in summary_rows:
    set_cell(ws, f"F{r}", label, font=F_LABEL, fill=FILL_SUMMARY, align=ALIGN_LEFT_NOWRAP)
    ws.merge_cells(f"G{r}:H{r}")
    is_big = label == "残り予算"
    set_cell(ws, f"G{r}", formula, font=(F_SUMMARY_BIG if is_big else F_CALC), fill=FILL_CALC,
              align=ALIGN_RIGHT, number_format=fmt)
    ws.row_dimensions[r].height = 22 if not is_big else 26

ws.conditional_formatting.add("G6:H6", Rule(
    type="expression", formula=["G6<0"], dxf=DifferentialStyle(font=Font(name=FONT_NAME, size=13, bold=True, color=COLOR_BAD))))
ws.conditional_formatting.add("G6:H6", Rule(
    type="expression", formula=["G6>=0"], dxf=DifferentialStyle(font=Font(name=FONT_NAME, size=13, bold=True, color=COLOR_GOOD))))

page_setup(ws, orientation="portrait", repeat_rows="1:2")

# ===========================================================================
# 5. 持ち物チェックリスト
# ===========================================================================
ws = wb.create_sheet("持ち物チェックリスト")
set_col_widths(ws, {"A": 12, "B": 24, "C": 10, "D": 14, "E": 26})
title_banner(ws, "E", "🎒 持ち物チェックリスト")
table_header(ws, 3, {"A": "カテゴリ", "B": "持ち物", "C": "必要数", "D": "準備済みチェック", "E": "メモ"})
ws.freeze_panes = "A4"

pack_examples = [
    ("衣類", "着替え(3日分)", 3, "✓", "現地は暑いので薄手のもの"),
    ("洗面用具", "歯ブラシセット", 4, "✓", ""),
    ("電子機器", "スマホ充電器", 2, "未", "モバイルバッテリーも忘れずに"),
    ("書類", "身分証明書・保険証", 4, "未", "免許証・保険証のコピー"),
    ("その他", "日焼け止め", 1, "未", ""),
]
DATA_START = 4
DATA_END = 28
for i, (cat, item, qty, chk, memo) in enumerate(pack_examples):
    r = DATA_START + i
    set_cell(ws, f"A{r}", cat, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_CENTER)
    set_cell(ws, f"B{r}", item, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP)
    set_cell(ws, f"C{r}", qty, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_CENTER)
    set_cell(ws, f"D{r}", chk, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_CENTER)
    set_cell(ws, f"E{r}", memo, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP)
    ws.row_dimensions[r].height = 20

for r in range(DATA_START + len(pack_examples), DATA_END + 1):
    for col, align in (("A", ALIGN_CENTER), ("B", ALIGN_LEFT_NOWRAP), ("C", ALIGN_CENTER),
                        ("D", ALIGN_CENTER), ("E", ALIGN_LEFT_NOWRAP)):
        set_cell(ws, f"{col}{r}", None, font=F_BODY, fill=FILL_INPUT, align=align)
    ws.row_dimensions[r].height = 20
    band_row(ws, r, ["A", "B", "C", "D", "E"], (r - DATA_START) % 2 == 1)

PROG_ROW = DATA_END + 2
ws.merge_cells(f"A{PROG_ROW}:B{PROG_ROW}")
set_cell(ws, f"A{PROG_ROW}", "準備進捗", font=F_TOTAL_LABEL, fill=FILL_SUMMARY, align=ALIGN_LEFT_NOWRAP)
set_cell(ws, f"C{PROG_ROW}", f'=COUNTA(B{DATA_START}:B{DATA_END})', font=F_CALC, fill=FILL_CALC,
          align=ALIGN_CENTER, border=BORDER_ALL)
set_cell(ws, f"D{PROG_ROW}", f'=COUNTIF(D{DATA_START}:D{DATA_END},"✓")', font=F_CALC, fill=FILL_CALC,
          align=ALIGN_CENTER, border=BORDER_ALL)
set_cell(ws, f"E{PROG_ROW}", f'=IFERROR(D{PROG_ROW}/C{PROG_ROW},0)', font=F_CALC, fill=FILL_CALC,
          align=ALIGN_CENTER, border=BORDER_ALL, number_format=PCT_FMT)
ws.row_dimensions[PROG_ROW].height = 22
NOTE_ROW = PROG_ROW + 1
ws.merge_cells(f"A{NOTE_ROW}:E{NOTE_ROW}")
set_cell(ws, f"A{NOTE_ROW}", "※準備進捗は「必要な持ち物の数」「✓済みの数」「達成率」の順に自動計算されます。", font=F_NOTE,
          fill=None, align=ALIGN_LEFT_NOWRAP, border=None)

add_list_validation(ws, f"D{DATA_START}:D{DATA_END}", "✓,未")
ws.conditional_formatting.add(f"D{DATA_START}:D{DATA_END}",
    CellIsRule(operator="equal", formula=['"✓"'], font=equal_rule(COLOR_GOOD, True)))
ws.conditional_formatting.add(f"B{DATA_START}:B{DATA_END}", Rule(
    type="expression", formula=[f'$D{DATA_START}="✓"'],
    dxf=DifferentialStyle(font=Font(name=FONT_NAME, size=10, strike=True, color="A6A6A6"))))

page_setup(ws, orientation="portrait")

# ===========================================================================
# 6. 行きたい場所・食べたいもの
# ===========================================================================
ws = wb.create_sheet("行きたい場所・食べたいもの")
set_col_widths(ws, {"A": 22, "B": 14, "C": 14, "D": 8, "E": 16, "F": 26, "G": 20})
title_banner(ws, "G", "📍 行きたい場所・食べたいもの")
table_header(ws, 3, {"A": "名称", "B": "カテゴリ", "C": "エリア", "D": "優先度", "E": "営業時間",
                       "F": "URL", "G": "メモ"})
ws.freeze_panes = "A4"

spot_examples = [
    ("首里城公園", "観光地", "那覇市", "高", "9:00-18:00", "https://oki-park.jp/shurijo/", "世界遺産"),
    ("国際通り", "ショッピング", "那覇市", "中", "10:00-22:00", "", "お土産・食べ歩き"),
    ("青の洞窟シュノーケリング", "体験・アクティビティ", "恩納村", "高", "要予約", "", "事前予約必須"),
    ("沖縄そば専門店", "グルメ", "恩納村", "中", "11:00-15:00", "", "地元で人気の店"),
]
DATA_START = 4
DATA_END = 28
for i, (name, cat, area, pri, hours, url, memo) in enumerate(spot_examples):
    r = DATA_START + i
    set_cell(ws, f"A{r}", name, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP)
    set_cell(ws, f"B{r}", cat, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_CENTER)
    set_cell(ws, f"C{r}", area, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_CENTER)
    set_cell(ws, f"D{r}", pri, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_CENTER)
    set_cell(ws, f"E{r}", hours, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_CENTER)
    set_cell(ws, f"F{r}", url, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP)
    set_cell(ws, f"G{r}", memo, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP)
    ws.row_dimensions[r].height = 20

for r in range(DATA_START + len(spot_examples), DATA_END + 1):
    for col, align in (("A", ALIGN_LEFT_NOWRAP), ("B", ALIGN_CENTER), ("C", ALIGN_CENTER),
                        ("D", ALIGN_CENTER), ("E", ALIGN_CENTER), ("F", ALIGN_LEFT_NOWRAP), ("G", ALIGN_LEFT_NOWRAP)):
        set_cell(ws, f"{col}{r}", None, font=F_BODY, fill=FILL_INPUT, align=align)
    ws.row_dimensions[r].height = 20
    band_row(ws, r, ["A", "B", "C", "D", "E", "F", "G"], (r - DATA_START) % 2 == 1)

add_list_validation(ws, f"B{DATA_START}:B{DATA_END}", "観光地,グルメ,カフェ,ショッピング,体験・アクティビティ,その他")
add_list_validation(ws, f"D{DATA_START}:D{DATA_END}", "高,中,低")
ws.conditional_formatting.add(f"D{DATA_START}:D{DATA_END}",
    CellIsRule(operator="equal", formula=['"高"'], font=equal_rule(COLOR_BAD, True)))
ws.conditional_formatting.add(f"D{DATA_START}:D{DATA_END}",
    CellIsRule(operator="equal", formula=['"中"'], font=equal_rule(COLOR_WARN, True)))
ws.conditional_formatting.add(f"D{DATA_START}:D{DATA_END}",
    CellIsRule(operator="equal", formula=['"低"'], font=equal_rule(COLOR_GOOD, False)))

page_setup(ws, orientation="landscape")

# ===========================================================================
# 7. お土産リスト
# ===========================================================================
ws = wb.create_sheet("お土産リスト")
set_col_widths(ws, {"A": 14, "B": 24, "C": 11, "D": 11, "E": 24})
title_banner(ws, "E", "🎁 お土産リスト")
table_header(ws, 3, {"A": "渡す相手", "B": "商品候補", "C": "予算", "D": "購入済み", "E": "メモ"})
ws.freeze_panes = "A4"

gift_examples = [
    ("職場の皆さん", "紅芋タルト(個包装)", 3000, "未", "20個入りを選ぶ"),
    ("友人Aさん", "シーサーの置物", 2000, "未", ""),
    ("家族用", "海ぶどう・ちんすこう", 4000, "済", ""),
]
DATA_START = 4
DATA_END = 23
for i, (whom, item, budget, bought, memo) in enumerate(gift_examples):
    r = DATA_START + i
    set_cell(ws, f"A{r}", whom, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP)
    set_cell(ws, f"B{r}", item, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP)
    set_cell(ws, f"C{r}", budget, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_RIGHT, number_format=CUR_FMT_SIMPLE)
    set_cell(ws, f"D{r}", bought, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_CENTER)
    set_cell(ws, f"E{r}", memo, font=F_EXAMPLE, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP)
    ws.row_dimensions[r].height = 20

for r in range(DATA_START + len(gift_examples), DATA_END + 1):
    for col, fmt, align in (("A", None, ALIGN_LEFT_NOWRAP), ("B", None, ALIGN_LEFT_NOWRAP),
                              ("C", CUR_FMT_SIMPLE, ALIGN_RIGHT), ("D", None, ALIGN_CENTER),
                              ("E", None, ALIGN_LEFT_NOWRAP)):
        set_cell(ws, f"{col}{r}", None, font=F_BODY, fill=FILL_INPUT, align=align, number_format=fmt)
    ws.row_dimensions[r].height = 20
    band_row(ws, r, ["A", "B", "C", "D", "E"], (r - DATA_START) % 2 == 1)

TOTAL_ROW = DATA_END + 1
ws.merge_cells(f"A{TOTAL_ROW}:B{TOTAL_ROW}")
set_cell(ws, f"A{TOTAL_ROW}", "合計予算", font=F_TOTAL_LABEL, fill=FILL_SUMMARY, align=ALIGN_LEFT_NOWRAP, border=BORDER_TOTAL)
set_cell(ws, f"C{TOTAL_ROW}", f"=SUM(C{DATA_START}:C{DATA_END})", font=F_CALC, fill=FILL_CALC,
          align=ALIGN_RIGHT, border=BORDER_TOTAL, number_format=CUR_FMT)
set_cell(ws, f"D{TOTAL_ROW}", "", fill=FILL_SUMMARY, border=BORDER_TOTAL)
set_cell(ws, f"E{TOTAL_ROW}", "", fill=FILL_SUMMARY, border=BORDER_TOTAL)

add_list_validation(ws, f"D{DATA_START}:D{DATA_END}", "済,未")
ws.conditional_formatting.add(f"D{DATA_START}:D{DATA_END}",
    CellIsRule(operator="equal", formula=['"済"'], font=equal_rule(COLOR_GOOD, True)))

page_setup(ws, orientation="portrait")

# ===========================================================================
# 8. 緊急連絡先
# ===========================================================================
ws = wb.create_sheet("緊急連絡先")
set_col_widths(ws, {"A": 20, "B": 15, "C": 18, "D": 28})
title_banner(ws, "D", "🚨 緊急連絡先")
table_header(ws, 3, {"A": "名称", "B": "電話番号", "C": "用途", "D": "メモ"})
ws.freeze_panes = "A4"

contact_examples = [
    ("警察", "110", "事件・事故", ""),
    ("消防・救急", "119", "火事・急病", ""),
    ("海上保安庁", "118", "海の事故", ""),
    ("宿泊先(〇〇リゾート&スパ)", "098-000-0000", "宿泊トラブル", "例:チェックイン遅延など"),
    ("旅行保険会社", "0120-000-000", "保険申請・事故対応", "契約番号を控えておく"),
]
DATA_START = 4
DATA_END = 18
for i, (name, tel, use, memo) in enumerate(contact_examples):
    r = DATA_START + i
    is_official = i < 3
    font_use = F_BODY if is_official else F_EXAMPLE
    set_cell(ws, f"A{r}", name, font=font_use, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP)
    set_cell(ws, f"B{r}", tel, font=font_use, fill=FILL_INPUT, align=ALIGN_CENTER)
    set_cell(ws, f"C{r}", use, font=font_use, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP)
    set_cell(ws, f"D{r}", memo, font=font_use, fill=FILL_INPUT, align=ALIGN_LEFT_NOWRAP)
    ws.row_dimensions[r].height = 20

for r in range(DATA_START + len(contact_examples), DATA_END + 1):
    for col, align in (("A", ALIGN_LEFT_NOWRAP), ("B", ALIGN_CENTER), ("C", ALIGN_LEFT_NOWRAP), ("D", ALIGN_LEFT_NOWRAP)):
        set_cell(ws, f"{col}{r}", None, font=F_BODY, fill=FILL_INPUT, align=align)
    ws.row_dimensions[r].height = 20
    band_row(ws, r, ["A", "B", "C", "D"], (r - DATA_START) % 2 == 1)

page_setup(ws, orientation="portrait")

# ---------------------------------------------------------------------------
# 仕上げ: 全シート共通のタブ色・アクティブシート
# ---------------------------------------------------------------------------
for sheet in wb.worksheets:
    sheet.sheet_properties.tabColor = COLOR_TITLE_BG
    sheet.sheet_view.zoomScale = 100

wb.active = 0
wb.save(OUTPUT_PATH)
print(f"saved: {OUTPUT_PATH}")
